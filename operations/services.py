from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO

from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import (
    IncomeItem,
    IndemnityLookupValue,
    IndemnityRecord,
    MeatCategory,
    MeatProductionEntry,
    ResidueCategory,
    ResidueCollection,
    WeeklyIncomeEntry,
)


@dataclass
class ImportResult:
    created_count: int
    errors: list[str]


def to_decimal(value, default="0"):
    if value in (None, ""):
        return Decimal(default)
    if isinstance(value, Decimal):
        return value
    normalized = str(value).strip().replace("R$", "").replace(".", "").replace(",", ".")
    return Decimal(normalized or default)


def parse_date(value):
    if hasattr(value, "strftime"):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Data inválida: {value}")


def normalize_headers(headers):
    return [str(header).strip().lower() if header is not None else "" for header in headers]


def load_rows_from_file(uploaded_file):
    suffix = uploaded_file.name.lower()
    if suffix.endswith(".csv"):
        content = uploaded_file.read().decode("utf-8-sig")
        reader = csv.DictReader(StringIO(content))
        return [dict(row) for row in reader]

    workbook = load_workbook(uploaded_file, data_only=True)
    sheet = workbook.active
    header_row = normalize_headers(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        rows.append(dict(zip(header_row, values)))
    return rows


def semaphore(color, label, detail):
    return {"color": color, "label": label, "detail": detail}


def income_semaphore(entry):
    expected_total = entry.quantity * entry.item.default_price
    if expected_total <= 0:
        return semaphore("secondary", "Sem base", "Item sem preço padrão configurado.")
    ratio = entry.calculated_total / expected_total
    if ratio >= Decimal("1"):
        return semaphore("success", "Acima da base", f"{ratio:.0%} da meta base.")
    if ratio >= Decimal("0.9"):
        return semaphore("warning", "Perto da base", f"{ratio:.0%} da meta base.")
    return semaphore("danger", "Abaixo da base", f"{ratio:.0%} da meta base.")


def meat_semaphore(entry):
    baseline = max(entry.average_general_grams, entry.average_previous_month_grams)
    if baseline <= 0:
        baseline = entry.average_per_head_grams
    if baseline <= 0:
        return semaphore("secondary", "Sem base", "Cadastre médias para análise.")
    ratio = entry.average_per_head_grams / baseline
    if ratio >= Decimal("1"):
        return semaphore("success", "Rendimento forte", f"{ratio:.0%} do benchmark.")
    if ratio >= Decimal("0.9"):
        return semaphore("warning", "Atenção", f"{ratio:.0%} do benchmark.")
    return semaphore("danger", "Queda relevante", f"{ratio:.0%} do benchmark.")


def residue_semaphore(entry, monthly_total):
    target = entry.category.target_amount
    if target <= 0:
        return semaphore("secondary", "Sem meta", "Categoria sem meta mensal.")
    ratio = monthly_total / target
    if ratio >= Decimal("1"):
        return semaphore("success", "Meta atingida", f"{ratio:.0%} da meta mensal.")
    if ratio >= Decimal("0.8"):
        return semaphore("warning", "Acompanhar", f"{ratio:.0%} da meta mensal.")
    return semaphore("danger", "Abaixo da meta", f"{ratio:.0%} da meta mensal.")


def indemnity_semaphore(entry):
    if entry.net_amount <= 0:
        return semaphore("success", "Revertida", "Sem impacto líquido.")
    if entry.net_amount <= Decimal("500"):
        return semaphore("warning", "Impacto moderado", f"R$ {entry.net_amount} líquido.")
    return semaphore("danger", "Impacto alto", f"R$ {entry.net_amount} líquido.")


def attach_income_status(entries):
    for entry in entries:
        entry.semaphore = income_semaphore(entry)
    return entries


def attach_meat_status(entries):
    for entry in entries:
        entry.semaphore = meat_semaphore(entry)
    return entries


def attach_residue_status(entries, queryset=None):
    queryset = queryset or ResidueCollection.objects.all()
    totals = {}
    for entry in entries:
        key = (entry.category_id, entry.collected_on.year, entry.collected_on.month)
        if key not in totals:
            totals[key] = (
                queryset.filter(
                    category_id=entry.category_id,
                    collected_on__year=entry.collected_on.year,
                    collected_on__month=entry.collected_on.month,
                ).aggregate(total=Sum("quantity"))["total"]
                or Decimal("0")
            )
        entry.monthly_total = totals[key]
        entry.semaphore = residue_semaphore(entry, entry.monthly_total)
    return entries


def attach_indemnity_status(entries):
    for entry in entries:
        entry.semaphore = indemnity_semaphore(entry)
    return entries


def sync_indemnity_lookup_values(organization, *, product="", owner_name="", responsible_name="", reason=""):
    mapping = {
        IndemnityLookupValue.FieldTypes.PRODUCT: product,
        IndemnityLookupValue.FieldTypes.OWNER: owner_name,
        IndemnityLookupValue.FieldTypes.RESPONSIBLE: responsible_name,
        IndemnityLookupValue.FieldTypes.REASON: reason,
    }
    for field_type, raw_value in mapping.items():
        value = str(raw_value or "").strip()
        if not value:
            continue
        IndemnityLookupValue.objects.get_or_create(
            organization=organization,
            field_type=field_type,
            value=value,
            defaults={"is_active": True},
        )


def summarize_semaphores(entries):
    summary = {"success": 0, "warning": 0, "danger": 0, "secondary": 0}
    for entry in entries:
        color = getattr(entry, "semaphore", {}).get("color", "secondary")
        summary[color] = summary.get(color, 0) + 1
    return summary


def _income_item_from_row(organization, row):
    name = str(row.get("item") or row.get("nome") or "").strip()
    unit_price = to_decimal(row.get("preco_unitario") or row.get("preço unitário") or row.get("unit_price") or "0")
    return IncomeItem.objects.get_or_create(
        organization=organization,
        name=name,
        defaults={"unit": "UNIT", "default_price": unit_price or Decimal("0")},
    )[0]


def import_income_rows(organization, uploaded_file):
    rows = load_rows_from_file(uploaded_file)
    created_count = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        try:
            item = _income_item_from_row(organization, row)
            WeeklyIncomeEntry.objects.create(
                organization=organization,
                item=item,
                reference_year=int(row.get("ano") or row.get("reference_year")),
                reference_month=int(row.get("mes") or row.get("mês") or row.get("reference_month")),
                week_number=int(row.get("semana") or row.get("week_number")),
                quantity=int(to_decimal(row.get("quantidade") or row.get("quantity"))),
                unit_price=to_decimal(row.get("preco_unitario") or row.get("preço unitário") or row.get("unit_price")),
                notes=str(row.get("observacoes") or row.get("observações") or row.get("notes") or "").strip(),
            )
            created_count += 1
        except Exception as exc:
            errors.append(f"Linha {index}: {exc}")
    return ImportResult(created_count, errors)


def _meat_category_from_row(organization, row):
    name = str(row.get("categoria") or row.get("category") or "").strip()
    price = to_decimal(row.get("preco_kg") or row.get("preço por kg") or row.get("price_per_kg") or "0")
    return MeatCategory.objects.get_or_create(
        organization=organization,
        name=name,
        defaults={"price_per_kg": price, "previous_price_per_kg": price},
    )[0]


def import_meat_rows(organization, uploaded_file):
    rows = load_rows_from_file(uploaded_file)
    created_count = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        try:
            category = _meat_category_from_row(organization, row)
            MeatProductionEntry.objects.create(
                organization=organization,
                category=category,
                reference_year=int(row.get("ano") or row.get("reference_year")),
                reference_month=int(row.get("mes") or row.get("mês") or row.get("reference_month")),
                week_number=int(row.get("semana") or row.get("week_number")),
                slaughter_quantity=int(row.get("quantidade_abatida") or row.get("slaughter_quantity")),
                weight_kg=to_decimal(row.get("peso_kg") or row.get("weight_kg")),
                average_per_head_grams=to_decimal(row.get("media_cabeca_g") or row.get("média por cabeça g") or row.get("average_per_head_grams")),
                average_general_grams=to_decimal(row.get("media_geral_g") or row.get("média geral g") or row.get("average_general_grams") or "0"),
                average_previous_month_grams=to_decimal(row.get("media_mes_anterior_g") or row.get("média do mês anterior g") or row.get("average_previous_month_grams") or "0"),
                donation_kg=to_decimal(row.get("doacao_kg") or row.get("doação kg") or row.get("donation_kg") or "0"),
                price_per_kg=to_decimal(row.get("preco_kg") or row.get("preço por kg") or row.get("price_per_kg")),
                notes=str(row.get("observacoes") or row.get("observações") or row.get("notes") or "").strip(),
            )
            created_count += 1
        except Exception as exc:
            errors.append(f"Linha {index}: {exc}")
    return ImportResult(created_count, errors)


def _residue_category_from_row(organization, row):
    name = str(row.get("categoria") or row.get("category") or "").strip()
    unit = str(row.get("unidade") or row.get("unit") or "UNIT").upper()
    unit_price = to_decimal(row.get("valor_unitario") or row.get("valor unitário") or row.get("unit_price") or "0")
    return ResidueCategory.objects.get_or_create(
        organization=organization,
        name=name,
        defaults={"unit": unit, "target_amount": Decimal("0"), "unit_price": unit_price},
    )[0]


def import_residue_rows(organization, uploaded_file):
    rows = load_rows_from_file(uploaded_file)
    created_count = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        try:
            category = _residue_category_from_row(organization, row)
            ResidueCollection.objects.create(
                organization=organization,
                category=category,
                collected_on=parse_date(row.get("data") or row.get("collected_on")),
                quantity=to_decimal(row.get("quantidade") or row.get("quantity")),
                notes=str(row.get("observacoes") or row.get("observações") or row.get("notes") or "").strip(),
            )
            created_count += 1
        except Exception as exc:
            errors.append(f"Linha {index}: {exc}")
    return ImportResult(created_count, errors)


def import_indemnity_rows(organization, uploaded_file):
    rows = load_rows_from_file(uploaded_file)
    created_count = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        try:
            record = IndemnityRecord.objects.create(
                organization=organization,
                occurred_on=parse_date(row.get("data") or row.get("ocorrencia") or row.get("occurred_on")),
                product=str(row.get("produto") or row.get("product") or "").strip(),
                quantity_description=str(row.get("quantidade_peso") or row.get("quantity_description") or "").strip(),
                outgoing_amount=to_decimal(row.get("saida") or row.get("valor_saida") or row.get("outgoing_amount") or "0"),
                reversal_amount=to_decimal(row.get("reversao") or row.get("reversão") or row.get("reversal_amount") or "0"),
                owner_name=str(row.get("dono") or row.get("owner_name") or "").strip(),
                responsible_name=str(row.get("responsavel") or row.get("responsável") or row.get("responsible_name") or "").strip(),
                reason=str(row.get("motivo") or row.get("reason") or "").strip(),
                notes=str(row.get("observacoes") or row.get("observações") or row.get("notes") or "").strip(),
            )
            sync_indemnity_lookup_values(
                organization,
                product=record.product,
                owner_name=record.owner_name,
                responsible_name=record.responsible_name,
                reason=record.reason,
            )
            created_count += 1
        except Exception as exc:
            errors.append(f"Linha {index}: {exc}")
    return ImportResult(created_count, errors)


def workbook_response(filename, title, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados"
    sheet.append([title])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def pdf_response(filename, title, headers, rows):
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=18, rightMargin=18)
    styles = getSampleStyleSheet()
    data = [[Paragraph(f"<b>{col}</b>", styles["BodyText"]) for col in headers]]
    for row in rows:
        data.append([Paragraph(str(value), styles["BodyText"]) for value in row])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#8A2E1E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CFC3B6")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor("#F7F1EA")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story = [Paragraph(f"<b>{title}</b>", styles["Title"]), Spacer(1, 12), table]
    document.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return response
